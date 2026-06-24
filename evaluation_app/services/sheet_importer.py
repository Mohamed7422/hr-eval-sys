from __future__ import annotations

import csv
import re
from io import TextIOWrapper
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
from datetime import datetime, date

try:
    import openpyxl
except ImportError:
    openpyxl = None

from django.db import transaction
from django.contrib.auth import get_user_model

from evaluation_app.models import (
    Company, Department, Employee, EmployeePlacement,
    ManagerialLevel, EmpStatus, JobType, BranchType,
)
from accounts.models import Role

User = get_user_model()

# ────────────────────────── Public API ──────────────────────────

REQUIRED_FIELDS = ["email", "name", "company_name", "managerial_level", "status", "join_date"]


def parse_sheet(request) -> List[Dict[str, Any]]:
    """Parse CSV or XLSX from request.FILES['file'] into list[dict]."""
    if "file" not in request.FILES:
        raise ValueError('Upload a file under the "file" key.')

    f = request.FILES["file"]
    suffix = Path(f.name).suffix.lower()

    if suffix == ".csv":
        text = TextIOWrapper(f.file, encoding="utf-8-sig", newline="")
        reader = csv.DictReader(text)
        rows = list(reader)
        if not rows:
            raise ValueError("CSV file is empty.")
        return rows

    if suffix in {".xlsx", ".xls"}:
        if openpyxl is None:
            raise ValueError("XLSX support requires 'openpyxl'. Install it or upload CSV.")
        wb = openpyxl.load_workbook(f, data_only=True)
        ws = wb.active
        headers = [
            str(c.value).strip() if c.value is not None else ""
            for c in next(ws.iter_rows(min_row=1, max_row=1))
        ]
        rows = []
        for r in ws.iter_rows(min_row=2, values_only=True):
            if all(v is None for v in r):
                continue
            rows.append({headers[i]: r[i] for i in range(len(headers))})
        if not rows:
            raise ValueError("XLSX sheet is empty.")
        return rows

    raise ValueError(f"Unsupported file type '{suffix}'. Upload CSV or XLSX.")


def import_from_sheet(
    rows: List[Dict[str, Any]],
    *,
    dry_run: bool = False,
) -> Dict[str, Any]:
    """
    Import employees from sheet rows.
    Companies and departments must already exist.
    """
    cleaned = _clean_rows(rows)

    # ── 1. Validate required fields ──
    errors = _validate_required(cleaned)
    if errors:
        return {"status": "invalid", "errors": errors}

    # ── 2. Resolve companies ──
    company_names = {r["company_name"] for r in cleaned}
    companies = {c.name: c for c in Company.objects.filter(name__in=company_names)}
    missing = company_names - set(companies.keys())
    if missing:
        errs = []
        for r in cleaned:
            if r["company_name"] in missing:
                errs.append({
                    "row": r["_row"],
                    "errors": {"company_name": [f'Company "{r["company_name"]}" does not exist.']},
                })
        return {"status": "invalid", "errors": errs}

    # ── 3. Resolve departments ──
    dept_keys = {
        (r["company_name"], r["department_name"])
        for r in cleaned if r["department_name"]
    }
    dept_map: Dict[Tuple[str, str], Department] = {}
    if dept_keys:
        depts = Department.objects.filter(
            company__name__in=[k[0] for k in dept_keys],
            name__in=[k[1] for k in dept_keys],
        ).select_related("company")
        dept_map = {(d.company.name, d.name): d for d in depts}

    dept_errors = []
    for r in cleaned:
        if r["department_name"]:
            key = (r["company_name"], r["department_name"])
            if key not in dept_map:
                dept_errors.append({
                    "row": r["_row"],
                    "errors": {"department_name": [
                        f'Department "{r["department_name"]}" not found in company "{r["company_name"]}".'
                    ]},
                })
    if dept_errors:
        return {"status": "invalid", "errors": dept_errors}

    # ── 4. Prefetch existing users / employees for upsert ──
    emails = {r["email"] for r in cleaned}
    users_by_email = {
        u.email: u
        for u in User.objects.filter(email__in=emails)
    }

    code_keys = {
        (r["company_name"], r["employee_code"])
        for r in cleaned if r["employee_code"]
    }
    employees_by_code: Dict[Tuple[str, str], Employee] = {}
    if code_keys:
        emps = Employee.objects.filter(
            company__name__in=[k[0] for k in code_keys],
            employee_code__in=[k[1] for k in code_keys],
        ).select_related("company", "user")
        employees_by_code = {
            (e.company.name if e.company else "", e.employee_code): e
            for e in emps
        }

    # ── 5. Build create / update lists ──
    taken_usernames = set(User.objects.values_list("username", flat=True))

    to_create: List[Dict[str, Any]] = []
    to_update: list[tuple] = []

    for r in cleaned:
        company = companies[r["company_name"]]
        dept = dept_map.get((r["company_name"], r["department_name"])) if r["department_name"] else None

        # Resolve username
        username = r["username"] or _generate_username(r["email"], r["name"])
        username = _ensure_unique_username(username, taken_usernames)

        # Upsert check
        user = users_by_email.get(r["email"])
        employee = None
        if user:
            employee = getattr(user, "employee_profile", None)
        if not employee and r["employee_code"]:
            employee = employees_by_code.get((r["company_name"], r["employee_code"]))
            if employee:
                user = employee.user

        row_data = {
            "_row": r["_row"],
            "username": username,
            "email": r["email"],
            "name": r["name"],
            "role": r["role"],
            "phone": r["phone"],
            "country_code": r["country_code"],
            "position": r["position"],
            "employee_code": r["employee_code"],
            "managerial_level": r["managerial_level"],
            "status": r["status"],
            "join_date": r["join_date"],
            "job_type": r["job_type"],
            "location": r["location"],
            "branch": r["branch"],
            "company": company,
            "department": dept,
        }

        if employee:
            to_update.append((user, employee, row_data))
        elif user:
            to_update.append((user, None, row_data))
        else:
            to_create.append(row_data)

    # ── 6. Dry run ──
    if dry_run:
        return {
            "status": "success",
            "dry_run": True,
            "validated_count": len(cleaned),
            "to_create": len(to_create),
            "to_update": len(to_update),
            "message": "Dry run successful. Data is ready to import.",
        }

    # ── 7. Commit ──
    created = 0
    updated = 0

    with transaction.atomic():
        for row_data in to_create:
            user_obj = User.objects.create_user(
                username=row_data["username"],
                email=row_data["email"],
                password="defaultpassword123",
                name=row_data["name"],
                role=row_data["role"] or Role.EMP,
                phone=row_data["phone"] or "",
                country_code=row_data["country_code"] or "",
                position=row_data["position"] or "",
                is_default_password=True,
            )
            emp = Employee.objects.create(
                user=user_obj,
                company=row_data["company"],
                employee_code=row_data["employee_code"] or "",
                managerial_level=row_data["managerial_level"],
                status=row_data["status"],
                join_date=row_data["join_date"],
                job_type=row_data["job_type"] or JobType.FULL_TIME,
                location=row_data["location"] or "",
                branch=row_data["branch"] or BranchType.OFFICE,
            )
            if row_data["department"]:
                EmployeePlacement.objects.create(
                    employee=emp,
                    company=row_data["company"],
                    department=row_data["department"],
                )
            created += 1

        for user_obj, employee, row_data in to_update:
            # Update user fields
            user_obj.name = row_data["name"]
            user_obj.phone = row_data["phone"] or user_obj.phone
            user_obj.country_code = row_data["country_code"] or user_obj.country_code
            user_obj.position = row_data["position"] or user_obj.position
            if row_data["role"]:
                user_obj.role = row_data["role"]
            user_obj.save()

            if employee:
                employee.employee_code = row_data["employee_code"] or employee.employee_code
                employee.managerial_level = row_data["managerial_level"]
                employee.status = row_data["status"]
                employee.join_date = row_data["join_date"]
                if row_data["job_type"]:
                    employee.job_type = row_data["job_type"]
                if row_data["location"]:
                    employee.location = row_data["location"]
                if row_data["branch"]:
                    employee.branch = row_data["branch"]
                employee.company = row_data["company"]
                employee.save()
            else:
                employee = Employee.objects.create(
                    user=user_obj,
                    company=row_data["company"],
                    employee_code=row_data["employee_code"] or "",
                    managerial_level=row_data["managerial_level"],
                    status=row_data["status"],
                    join_date=row_data["join_date"],
                    job_type=row_data["job_type"] or JobType.FULL_TIME,
                    location=row_data["location"] or "",
                    branch=row_data["branch"] or BranchType.OFFICE,
                )

            # Upsert placement
            if row_data["department"]:
                EmployeePlacement.objects.update_or_create(
                    employee=employee,
                    defaults={
                        "company": row_data["company"],
                        "department": row_data["department"],
                    },
                )
            updated += 1

    return {"status": "imported", "created": created, "updated": updated}


# ────────────────────────── Helpers ──────────────────────────

def _clean_rows(rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    cleaned = []
    for i, raw in enumerate(rows, start=1):
        r = {k.strip(): v for k, v in raw.items()}
        row = {
            "_row": i,
            "employee_code": _str(r.get("employee_code")),
            "username": _str(r.get("username")),
            "email": _str(r.get("email"), lower=True),
            "role": _norm_choice(_str(r.get("role")), Role.choices, ROLE_ALIASES),
            "name": _str(r.get("name")),
            "company_name": _str(r.get("company_name")),
            "managerial_level": _norm_choice(
                _str(r.get("managerial_level")), ManagerialLevel.choices, MANAGERIAL_ALIASES
            ),
            "status": _norm_choice(_str(r.get("status")), EmpStatus.choices, STATUS_ALIASES),
            "join_date": _parse_date(r.get("join_date")),
            "country_code": _str(r.get("country_code")),
            "phone": _str(r.get("phone")),
            "position": _str(r.get("position")),
            "department_name": _str(r.get("department_name")),
            "job_type": _norm_choice(_str(r.get("job_type")), JobType.choices, JOBTYPE_ALIASES),
            "location": _str(r.get("location")),
            "branch": _norm_choice(_str(r.get("branch")), BranchType.choices, BRANCH_ALIASES),
        }
        cleaned.append(row)
    return cleaned


def _validate_required(cleaned):
    errors = []
    for row in cleaned:
        local = {}
        for field in REQUIRED_FIELDS:
            if not row.get(field):
                local[field] = ["This field is required."]
        if local:
            errors.append({"row": row["_row"], "errors": local})
    return errors


def _str(value, lower=False) -> Optional[str]:
    if value is None or (isinstance(value, str) and value.strip() == ""):
        return None
    s = str(value).strip()
    return s.lower() if lower else s


def _parse_date(value) -> Optional[str]:
    if value is None or value == "":
        return None
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    s = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return s


def _norm_key(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(s).lower())


def _norm_choice(value, choices, aliases=None):
    if not value:
        return None
    key = _norm_key(value)
    table = {}
    for stored, label in choices:
        table[_norm_key(stored)] = stored
        table[_norm_key(label)] = stored
    if aliases:
        for alias, real in aliases.items():
            table[_norm_key(alias)] = real
    return table.get(key, value)


def _generate_username(email: Optional[str], name: Optional[str]) -> str:
    if email:
        base = email.split("@")[0][:150]
    elif name:
        base = name.lower().replace(" ", ".")[:150]
    else:
        base = "user"
    return re.sub(r"[^a-zA-Z0-9_.]+", "", base) or "user"


def _ensure_unique_username(username: str, taken: set) -> str:
    candidate = username
    n = 1
    while candidate in taken:
        n += 1
        candidate = f"{username}{n}"
    taken.add(candidate)
    return candidate


# ────────────────────────── Choice Aliases ──────────────────────────

ROLE_ALIASES = {
    "admin": Role.ADMIN,
    "hr": Role.HR,
    "hod": Role.HOD,
    "head of dept": Role.HOD,
    "head of department": Role.HOD,
    "lm": Role.LM,
    "line manager": Role.LM,
    "emp": Role.EMP,
    "employee": Role.EMP,
}

MANAGERIAL_ALIASES = {
    "ic": ManagerialLevel.IC,
    "individual contributor": ManagerialLevel.IC,
    "supervisory": ManagerialLevel.SUPERVISORY,
    "middle": ManagerialLevel.MIDDLE,
    "middle management": ManagerialLevel.MIDDLE,
    "executive": ManagerialLevel.EXECUTIVE,
    "executive management": ManagerialLevel.EXECUTIVE,
}

STATUS_ALIASES = {
    "active": EmpStatus.ACTIVE,
    "inactive": EmpStatus.INACTIVE,
    "default active": EmpStatus.DEFAULT,
    "defaultactive": EmpStatus.DEFAULT,
}

JOBTYPE_ALIASES = {
    "full time": JobType.FULL_TIME,
    "fulltime": JobType.FULL_TIME,
    "full-time": JobType.FULL_TIME,
    "part time": JobType.PART_TIME,
    "parttime": JobType.PART_TIME,
    "part-time": JobType.PART_TIME,
    "full time remote": JobType.FULL_TIME_REMOTE,
    "full-time remote": JobType.FULL_TIME_REMOTE,
    "part time remote": JobType.PART_TIME_REMOTE,
    "part-time remote": JobType.PART_TIME_REMOTE,
}

BRANCH_ALIASES = {
    "office": BranchType.OFFICE,
    "store": BranchType.STORE,
}
